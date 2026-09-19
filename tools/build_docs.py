"""Build .docx and .pdf versions of a Markdown document.

    python tools/build_docs.py docs/RAD-SMART_Research_Report.md

Writes <name>.docx (editable, python-docx) and <name>.pdf (print layout via
headless Chrome/Edge) next to the Markdown file. Supports the Markdown subset
used in this repo: #-headings, paragraphs, **bold**, *italic*, `code`, links,
bullet / numbered lists (2 levels), pipe tables, images with captions,
> call-outs, fenced code, --- rules and <!-- pagebreak -->.
"""

from __future__ import annotations

import html
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu

ACCENT = "1C5CAB"
INK2 = "52514E"
BAND = "F4F3EF"

# ------------------------------------------------------------------ parsing
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def inline_runs(text):
    """-> list of (text, style) where style in {'', 'b', 'i', 'code', ('link', url)}."""
    out = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], "b"))
        elif part.startswith("`") and part.endswith("`"):
            out.append((part[1:-1], "code"))
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            out.append((m.group(1), ("link", m.group(2))))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], "i"))
        else:
            out.append((part, ""))
    return out


def parse(md: str):
    lines = md.splitlines()
    blocks, i = [], 0
    para = []

    def flush():
        if para:
            blocks.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("<!-- pagebreak -->"):
            flush(); blocks.append(("pagebreak",)); i += 1; continue
        if s.startswith("<!--"):
            flush()
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1; continue
        if not s:
            flush(); i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush(); blocks.append(("h", len(m.group(1)), m.group(2).strip())); i += 1; continue
        if s.startswith("```"):
            flush(); code = []; i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            blocks.append(("code", "\n".join(code))); i += 1; continue
        if s == "---":
            flush(); blocks.append(("hr",)); i += 1; continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)\})?$", s)
        if m:
            flush(); blocks.append(("img", m.group(2), m.group(1), float(m.group(4) or 0))); i += 1; continue
        if s.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows[0], rows[1:])); continue
        if s.startswith(">"):
            flush(); q = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q.append(lines[i].strip()[1:].strip()); i += 1
            blocks.append(("quote", " ".join(q))); continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            flush(); items = []; kind = "ol" if m.group(2)[0].isdigit() else "ul"
            while i < len(lines):
                m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if m:
                    level = 1 if len(m.group(1).replace("\t", "    ")) >= 2 else 0
                    items.append([level, m.group(2), m.group(3).strip()]); i += 1
                elif lines[i].strip() and lines[i].startswith("  ") and items:
                    items[-1][2] += " " + lines[i].strip(); i += 1
                else:
                    break
            blocks.append((kind, items)); continue
        para.append(line); i += 1
    flush()
    return blocks


# --------------------------------------------------------------------- DOCX
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _para_border(p, side="bottom", color="C3C2B7", size=6, space=4):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    bdr.append(el); pPr.append(bdr)


def _para_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _hyperlink(p, text, url):
    part = p.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), ACCENT); rPr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    r.append(rPr); t = OxmlElement("w:t"); t.text = text
    t.set(qn("xml:space"), "preserve"); r.append(t); link.append(r); p._p.append(link)


def _add_inline(p, text, size=None, color=None, bold=False):
    for chunk, style in inline_runs(text):
        if isinstance(style, tuple):
            _hyperlink(p, chunk, style[1]); continue
        r = p.add_run(chunk)
        r.bold = bold or style == "b"
        r.italic = style == "i"
        if style == "code":
            r.font.name = "Consolas"; r.font.size = Pt(9)
        elif size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def _page_number_footer(section, label):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(label + "   ·   page "); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(INK2)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run(); run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(INK2)
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt
            run._r.append(it)


def to_docx(blocks, out: Path, base: Path, footer_label: str):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    width_in = (21.0 - 4.0) / 2.54
    st = doc.styles
    st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(10.5)
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st["Normal"].paragraph_format.space_after = Pt(5)
    st["Normal"].paragraph_format.line_spacing = 1.12
    for name, size, before in (("Title", 24, 0), ("Heading 1", 16, 14), ("Heading 2", 12.5, 10),
                               ("Heading 3", 11, 8)):
        s = st[name]; s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(ACCENT if name != "Title" else "0B0B0B")
        rf = s.element.rPr.find(qn("w:rFonts"))
        if rf is not None:
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            rf.set(qn("w:ascii"), "Calibri"); rf.set(qn("w:hAnsi"), "Calibri")
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(4)
    _page_number_footer(sec, footer_label)

    for b in blocks:
        kind = b[0]
        if kind == "h":
            level, text = b[1], b[2]
            if level == 1:
                p = doc.add_paragraph(style="Title"); _add_inline(p, text)
            else:
                p = doc.add_paragraph(style=f"Heading {level - 1}")
                _add_inline(p, text)
                p.paragraph_format.keep_with_next = True
        elif kind == "p":
            p = doc.add_paragraph(); _add_inline(p, b[1])
        elif kind in ("ul", "ol"):
            counters = [0, 0]
            for level, marker, text in b[1]:
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.left_indent = Cm(0.6 + 0.6 * level); pf.first_line_indent = Cm(-0.45)
                pf.space_after = Pt(2)
                if kind == "ol" and marker[0].isdigit():
                    counters[level] += 1
                    if level == 0:
                        counters[1] = 0
                    bullet = f"{counters[level]}."
                else:
                    bullet = "•" if level == 0 else "–"
                r = p.add_run(bullet + "\t"); r.font.color.rgb = RGBColor.from_string(ACCENT)
                pf.tab_stops.add_tab_stop(Cm(0.6 + 0.6 * level))
                _add_inline(p, text)
        elif kind == "table":
            header, rows = b[1], b[2]
            ncol = len(header)
            t = doc.add_table(rows=1 + len(rows), cols=ncol)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.style = doc.styles["Table Grid"]
            lens = [max([len(re.sub(r"[*`]", "", header[c]))] +
                        [len(re.sub(r"[*`]", "", r[c])) if c < len(r) else 0 for r in rows]) for c in range(ncol)]
            weights = [min(max(l, 6), 60) ** 0.8 for l in lens]
            widths = [width_in * w / sum(weights) for w in weights]
            for ri, row in enumerate([header] + rows):
                for c in range(ncol):
                    cell = t.cell(ri, c)
                    cell.width = Emu(int(widths[c] * 914400))
                    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                    txt = row[c] if c < len(row) else ""
                    _add_inline(cell.paragraphs[0], txt, size=9, bold=(ri == 0))
                    if ri == 0:
                        _shade(cell, "E8EFF9")
            _set_table_borders(t)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "img":
            path, caption, w = (base / b[1]).resolve(), b[2], b[3]
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Cm(w * 2.54) if w else Cm(17.0))
            p.paragraph_format.keep_with_next = True
            if caption:
                c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _add_inline(c, caption, size=9, color=INK2)
                c.runs[0].italic = False
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.3)
            _para_shade(p, BAND); _para_border(p, "left", ACCENT, 18, 8)
            _add_inline(p, b[1])
        elif kind == "code":
            for ln in b[1].split("\n"):
                p = doc.add_paragraph(); _para_shade(p, BAND)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(8.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif kind == "hr":
            p = doc.add_paragraph(); _para_border(p)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(out)


def _set_table_borders(t):
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D5D3CC")
        borders.append(el)
    tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 40), ("bottom", 40), ("left", 80), ("right", 80)):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)


# --------------------------------------------------------------------- HTML
CSS = """
@page { size: A4; margin: 16mm 16mm 18mm 16mm; }
:root { --ink:#0b0b0b; --ink2:#52514e; --accent:#1c5cab; --band:#f4f3ef; --hair:#d5d3cc; }
* { box-sizing: border-box; }
body { font: 10pt/1.45 "Segoe UI", Calibri, system-ui, sans-serif; color: var(--ink); margin: 0; }
h1 { font-size: 22pt; margin: 0 0 6pt; line-height: 1.15; }
h2 { font-size: 15pt; color: var(--accent); margin: 18pt 0 6pt; break-after: avoid; }
h3 { font-size: 12pt; color: var(--accent); margin: 12pt 0 4pt; break-after: avoid; }
h4 { font-size: 10.5pt; margin: 10pt 0 3pt; break-after: avoid; }
p { margin: 0 0 6pt; }
ul, ol { margin: 0 0 6pt; padding-left: 16pt; }
li { margin: 1.5pt 0; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt; font-size: 8.6pt; break-inside: auto; }
tr { break-inside: avoid; }
th { background: #e8eff9; text-align: left; font-weight: 600; }
th, td { border: 0.6pt solid var(--hair); padding: 3pt 5pt; vertical-align: top; }
code { font: 8.6pt Consolas, monospace; background: var(--band); padding: 0 2pt; border-radius: 2pt; }
pre { font: 8pt/1.35 Consolas, monospace; background: var(--band); padding: 7pt 9pt; border-radius: 4pt;
      white-space: pre-wrap; break-inside: avoid; }
blockquote { margin: 6pt 0 8pt; padding: 6pt 10pt; background: var(--band); border-left: 3pt solid var(--accent);
             border-radius: 0 4pt 4pt 0; }
figure { margin: 6pt 0 10pt; break-inside: avoid; }
figure img { width: 100%; display: block; }
figcaption { font-size: 8.5pt; color: var(--ink2); margin-top: 3pt; }
hr { border: 0; border-top: 0.8pt solid var(--hair); margin: 10pt 0; }
a { color: var(--accent); text-decoration: none; }
.pb { break-after: page; }
"""


def _inline_html(text):
    out = []
    for chunk, style in inline_runs(text):
        c = html.escape(chunk)
        if isinstance(style, tuple):
            out.append(f'<a href="{html.escape(style[1])}">{c}</a>')
        elif style == "b":
            out.append(f"<strong>{c}</strong>")
        elif style == "i":
            out.append(f"<em>{c}</em>")
        elif style == "code":
            out.append(f"<code>{c}</code>")
        else:
            out.append(c)
    return "".join(out)


def to_html(blocks, base: Path, title: str) -> str:
    h = [f"<!DOCTYPE html><html lang='en'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
         f"<style>{CSS}</style></head><body>"]
    for b in blocks:
        k = b[0]
        if k == "h":
            h.append(f"<h{b[1]}>{_inline_html(b[2])}</h{b[1]}>")
        elif k == "p":
            h.append(f"<p>{_inline_html(b[1])}</p>")
        elif k in ("ul", "ol"):
            tag = "ol" if k == "ol" else "ul"
            h.append(f"<{tag}>")
            open_sub = False
            for level, _, text in b[1]:
                if level == 1 and not open_sub:
                    h.append("<ul>"); open_sub = True
                elif level == 0 and open_sub:
                    h.append("</ul>"); open_sub = False
                h.append(f"<li>{_inline_html(text)}</li>")
            if open_sub:
                h.append("</ul>")
            h.append(f"</{tag}>")
        elif k == "table":
            h.append("<table><thead><tr>" + "".join(f"<th>{_inline_html(c)}</th>" for c in b[1]) +
                     "</tr></thead><tbody>")
            for r in b[2]:
                h.append("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in r) + "</tr>")
            h.append("</tbody></table>")
        elif k == "img":
            src = (base / b[1]).resolve().as_uri()
            style = f' style="width:{b[3] / 6.7 * 100:.0f}%"' if b[3] else ""
            h.append(f'<figure><img src="{src}"{style}><figcaption>{_inline_html(b[2])}</figcaption></figure>')
        elif k == "quote":
            h.append(f"<blockquote>{_inline_html(b[1])}</blockquote>")
        elif k == "code":
            h.append(f"<pre>{html.escape(b[1])}</pre>")
        elif k == "hr":
            h.append("<hr>")
        elif k == "pagebreak":
            h.append('<div class="pb"></div>')
    h.append("</body></html>")
    return "\n".join(h)


def find_browser():
    for p in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
              shutil.which("chrome"), shutil.which("msedge"), shutil.which("chromium")):
        if p and Path(p).exists():
            return p
    return None


def main(md_path: str):
    md_file = Path(md_path).resolve()
    blocks = parse(md_file.read_text(encoding="utf-8"))
    title = next((b[2] for b in blocks if b[0] == "h" and b[1] == 1), md_file.stem)
    docx_out = md_file.with_suffix(".docx")
    to_docx(blocks, docx_out, md_file.parent, "RAD-SMART · Health-a-thon 2026")
    print("wrote", docx_out)
    html_out = md_file.with_suffix(".html")
    html_out.write_text(to_html(blocks, md_file.parent, title), encoding="utf-8")
    browser = find_browser()
    if browser:
        pdf_out = md_file.with_suffix(".pdf")
        subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_out}", html_out.as_uri()],
                       check=False, capture_output=True, timeout=180)
        print("wrote", pdf_out if pdf_out.exists() else "(pdf failed)")


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        main(arg)
